"""Validate actual browser downloads; requires python-docx for an editable round trip."""
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
import json, tempfile
from docx import Document
files=sorted(Path('artifacts/pr346').glob('memo-*.docx'))
assert len(files)==30, len(files)
results=[]
for path in files:
    with ZipFile(path) as z:
        assert z.testzip() is None
        for name in z.namelist():
            if name.endswith(('.xml','.rels')): ET.fromstring(z.read(name))
        assert {'[Content_Types].xml','_rels/.rels','word/document.xml','word/styles.xml'}.issubset(z.namelist())
        content=z.read('word/document.xml').decode()
        assert 'documentProtection' not in content
        for heading in ['บันทึกข้อความ','ส่วนราชการ','ที่','วันที่','เรื่อง','เรียน','เรื่องเดิม','ข้อเท็จจริง','ข้อพิจารณา','ข้อเสนอ','[รอตรวจสอบ]']: assert heading in content
    doc=Document(path)
    with tempfile.TemporaryDirectory() as directory:
        target=Path(directory)/'editable.docx'
        doc.add_paragraph('ข้อความแก้ไขทดสอบ')
        doc.save(target)
        assert Document(target).paragraphs[-1].text=='ข้อความแก้ไขทดสอบ'
    results.append({'file':path.name,'zip':'PASS','ooxml':'PASS','thaiUnicode':'PASS','editableRoundTrip':'PASS','microsoftWord':'NOT VERIFIED'})
Path('docs/quality/pr346/docx-results.json').write_text(json.dumps(results,ensure_ascii=False,indent=2))
print(f'PASS {len(results)} ZIP/OOXML/Thai/editable round trips; Microsoft Word NOT VERIFIED')
